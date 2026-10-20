from __future__ import annotations

import ast
import builtins
import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from metabatch.excel_processor import ExcelProcessingError, inspect_enrichment_sheet
from metabatch.gene_reader import (
    MESSAGE_TOO_MANY_GENES_REASON,
    SKIPPED_REASON_EXISTING,
    exceeds_gene_limit,
    too_many_genes_message,
)
from metabatch.i18n import UI_LANGUAGE_EN, UI_LANGUAGE_ZH, set_ui_language, tr
from metabatch.models import DEFAULT_MAX_GENES_PER_FILE, AppConfig, DownloadType, TaskResult, TaskStatus
from metabatch.paths import build_output_paths
from metabatch.workflow import (
    TRANSLATION_STATUS_PAUSED,
    GeneFileRecord,
    WorkflowController,
    _is_complete_result_workbook,
    _metascape_state,
    _translation_state,
    scan_gene_files,
)


def write_gene_file(path: Path, count: int) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(f"GENE{i}" for i in range(count)), encoding="utf-8")
    return path


def write_result_workbook(path: Path, headers: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Enrichment"
    sheet.append(headers)
    sheet.append(["1", "GO", "term", "description"])
    workbook.save(path)
    return path


def _binding_names(node: ast.AST) -> set[str]:
    """收集一个作用域里所有"被绑定"的名字：参数、赋值、导入、except as、with as、def/class。"""
    names: set[str] = set()
    for sub in ast.walk(node):
        if isinstance(sub, ast.Name) and isinstance(sub.ctx, (ast.Store, ast.Del)):
            names.add(sub.id)
        elif isinstance(sub, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            names.add(sub.name)
        elif isinstance(sub, (ast.Import, ast.ImportFrom)):
            for alias in sub.names:
                names.add(alias.asname or alias.name.split(".")[0])
        elif isinstance(sub, ast.arg):
            names.add(sub.arg)
        elif isinstance(sub, ast.ExceptHandler) and sub.name:
            names.add(sub.name)
        elif isinstance(sub, ast.Global | ast.Nonlocal):
            names.update(sub.names)
    return names


def _iter_own_code(node: ast.AST):
    """只遍历当前作用域自身的语句，不进入嵌套函数 / lambda / 推导式（它们各有自己的作用域）。"""

    def walk(current: ast.AST):
        yield current
        for child in ast.iter_child_nodes(current):
            if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef, ast.Lambda, ast.ListComp, ast.SetComp, ast.DictComp, ast.GeneratorExp)):
                continue
            yield from walk(child)

    yield from walk(node)


def _module_level_names(tree: ast.Module) -> set[str]:
    """只收集模块顶层绑定的名字（不进入函数/类体）。

    实际教训：之前用 ast.walk 收集"全局名"，run() 里的局部变量 metascape_connection_mode
    被当成全局，_process_single_file() 里误用它时守卫没报，真机运行全部文件失败。
    """
    names: set[str] = set()
    for stmt in tree.body:
        if isinstance(stmt, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            names.add(stmt.name)
            continue
        names |= _binding_names(stmt)
    return names


def _check_module(path: Path) -> list[str]:
    tree = ast.parse(path.read_text(encoding="utf-8"))
    globals_and_builtins = _module_level_names(tree) | set(dir(builtins)) | {"__name__", "__file__", "__doc__"}
    problems: list[str] = []

    def check(scope: ast.AST, available: set[str], label: str, outer: set[str]) -> None:
        local = _binding_names(scope) | outer
        for node in _iter_own_code(scope):
            if isinstance(node, ast.Name) and isinstance(node.ctx, ast.Load) and node.id not in local | available:
                problems.append(f"{path.name}:{node.lineno} {label} 使用了未定义的名字 {node.id}")
            elif isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr == "evaluate":
                for argument in node.args[1:]:
                    for sub in _iter_own_code(argument):
                        if isinstance(sub, ast.Name) and isinstance(sub.ctx, ast.Load) and sub.id not in local | available:
                            problems.append(
                                f"{path.name}:{argument.lineno} {label} 的 page.evaluate 参数使用了未定义的名字 {sub.id}"
                            )

    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            check(node, globals_and_builtins, f"{node.name}()", set())
        elif isinstance(node, ast.ClassDef):
            # 类方法也要查：方法体看不到类体里的名字，只能看到模块全局
            for item in node.body:
                if isinstance(item, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    check(item, globals_and_builtins, f"{node.name}.{item.name}()", set())
    return problems


class GeneLimitTests(unittest.TestCase):
    def test_default_limit_is_5000(self) -> None:
        config = AppConfig(input_dir=Path("."), output_dir=Path("."))
        self.assertTrue(config.skip_large_gene_files)
        self.assertEqual(config.max_genes_per_file, DEFAULT_MAX_GENES_PER_FILE)
        self.assertEqual(config.normalized_max_genes_per_file(), 5000)

    def test_disabling_the_switch_removes_the_limit(self) -> None:
        config = AppConfig(input_dir=Path("."), output_dir=Path("."), skip_large_gene_files=False)
        self.assertIsNone(config.normalized_max_genes_per_file())

    def test_zero_or_negative_limit_is_treated_as_unlimited(self) -> None:
        config = AppConfig(input_dir=Path("."), output_dir=Path("."), max_genes_per_file=0)
        self.assertIsNone(config.normalized_max_genes_per_file())

    def test_exceeds_gene_limit_boundaries(self) -> None:
        self.assertFalse(exceeds_gene_limit(5000, 5000))
        self.assertTrue(exceeds_gene_limit(5001, 5000))
        self.assertFalse(exceeds_gene_limit(99999, None))

    def test_message_mentions_both_numbers(self) -> None:
        message = too_many_genes_message(8000, 5000)
        self.assertIn("8000", message)
        self.assertIn("5000", message)


class ExistingResultDetectionTests(unittest.TestCase):
    def test_workbook_without_target_column_is_incomplete_when_translating(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = write_result_workbook(Path(tmp) / "a_metascape.xlsx", ["Category", "Description"])
            self.assertFalse(_is_complete_result_workbook(path, "中文描述"))
            self.assertTrue(_is_complete_result_workbook(path, None))

    def test_workbook_with_target_column_is_complete(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = write_result_workbook(Path(tmp) / "a_metascape.xlsx", ["Category", "Description", "中文描述"])
            self.assertTrue(_is_complete_result_workbook(path, "中文描述"))

    def test_workbook_without_enrichment_sheet_is_incomplete(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a_metascape.xlsx"
            workbook = Workbook()
            workbook.active.title = "Annotation"
            workbook.save(path)
            self.assertFalse(_is_complete_result_workbook(path, None))

    def test_corrupted_file_is_incomplete(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "broken_metascape.xlsx"
            path.write_bytes(b"not a workbook")
            self.assertFalse(_is_complete_result_workbook(path, None))

    def test_inspect_enrichment_sheet_reports_row_count_and_existing_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = write_result_workbook(Path(tmp) / "a_metascape.xlsx", ["Category", "Description"])
            self.assertEqual(inspect_enrichment_sheet(path), 1)

            translated = write_result_workbook(
                Path(tmp) / "b_metascape.xlsx", ["Category", "Description", "中文描述"]
            )
            self.assertEqual(inspect_enrichment_sheet(translated), 0)

    def test_inspect_enrichment_sheet_rejects_missing_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a_metascape.xlsx"
            workbook = Workbook()
            workbook.active.title = "Annotation"
            workbook.save(path)
            with self.assertRaises(ExcelProcessingError):
                inspect_enrichment_sheet(path)


class ScanGeneFilesTests(unittest.TestCase):
    """预扫描 = 续跑判定 + 阈值跳过 + 基因列表只读一次。"""

    def _build_tree(self, root: Path) -> tuple[Path, Path]:
        input_dir = root / "genes"
        output_dir = root / "out"
        input_dir.mkdir(parents=True)
        output_dir.mkdir(parents=True)
        return input_dir, output_dir

    def _config(self, input_dir: Path, output_dir: Path, **kwargs) -> AppConfig:
        return AppConfig(input_dir=input_dir, output_dir=output_dir, download_type=DownloadType.XLSX_ONLY, **kwargs)

    def test_buckets_pending_existing_and_oversized(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_dir, output_dir = self._build_tree(root)
            config = self._config(input_dir, output_dir, skip_large_gene_files=True, max_genes_per_file=5000)

            small = write_gene_file(input_dir / "small.txt", 10)
            big = write_gene_file(input_dir / "big.txt", 8000)
            done = write_gene_file(input_dir / "done.txt", 10)
            stale = write_gene_file(input_dir / "stale.txt", 10)
            pending_untranslated = write_gene_file(input_dir / "pending.txt", 10)

            # done 有 Enrichment + 中文描述 → 已完成
            write_result_workbook(
                build_output_paths(input_dir, output_dir, done, DownloadType.XLSX_ONLY).download_path,
                ["Category", "Description", "中文描述"],
            )
            # stale 有 Enrichment 但缺中文描述 → 需重跑
            write_result_workbook(
                build_output_paths(input_dir, output_dir, stale, DownloadType.XLSX_ONLY).download_path,
                ["Category", "Description"],
            )
            # pending 的结果文件损坏 → 需重跑
            broken = build_output_paths(input_dir, output_dir, pending_untranslated, DownloadType.XLSX_ONLY).download_path
            broken.parent.mkdir(parents=True, exist_ok=True)
            broken.write_bytes(b"broken")

            plan, annotated = scan_gene_files(config, [small, big, done, stale, pending_untranslated], "中文描述")

            # 文件顺序由 discover_gene_files 决定，这里只比较集合
            self.assertEqual(set(plan.pending), {small, stale, pending_untranslated})
            self.assertEqual([item[0] for item in plan.existing], [done])
            self.assertEqual([item[0] for item in plan.oversized], [big])
            self.assertEqual(len(annotated), 5)
            self.assertEqual(plan.gene_lists[small][:2], ["GENE0", "GENE1"])

    def test_limit_disabled_keeps_large_lists_pending(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_dir, output_dir = self._build_tree(root)
            config = self._config(input_dir, output_dir, skip_large_gene_files=False)
            big = write_gene_file(input_dir / "big.txt", 8000)

            plan, _ = scan_gene_files(config, [big], "中文描述")

            self.assertEqual(plan.pending, [big])
            self.assertEqual(plan.oversized, [])

    def test_translation_disabled_scores_completion_on_enrichment_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_dir, output_dir = self._build_tree(root)
            config = self._config(input_dir, output_dir, enable_translation=False)
            gene = write_gene_file(input_dir / "a.txt", 5)
            write_result_workbook(
                build_output_paths(input_dir, output_dir, gene, DownloadType.XLSX_ONLY).download_path,
                ["Category", "Description"],
            )

            plan, _ = scan_gene_files(config, [gene], "中文描述")

            self.assertEqual(plan.pending, [])
            self.assertEqual(len(plan.existing), 1)

    def test_unreadable_file_is_reported_separately(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_dir, output_dir = self._build_tree(root)
            config = self._config(input_dir, output_dir)
            broken = input_dir / "broken.xlsx"
            broken.write_bytes(b"not a workbook")

            plan, _ = scan_gene_files(config, [broken], "中文描述")

            self.assertEqual(plan.unreadable and plan.unreadable[0][0], broken)
            self.assertEqual(plan.pending, [])


class TranslationPreflightTests(unittest.TestCase):
    class _ExplodingTranslator:
        def __init__(self) -> None:
            self.calls = 0

        def translate(self, text: str) -> str:
            self.calls += 1
            raise RuntimeError("api.bblabu.ai: Read timed out")

    class _WorkingTranslator:
        def translate(self, text: str) -> str:
            return "细胞周期"

    def test_failing_preflight_pauses_translation(self) -> None:
        translator = self._ExplodingTranslator()
        logs: list[str] = []
        ok = WorkflowController._preflight_translation(
            translator, log_callback=lambda worker_id, message: logs.append(message)
        )
        self.assertFalse(ok)
        self.assertEqual(translator.calls, 1)  # 每次运行只做一次预检
        self.assertTrue(any("预检失败" in line for line in logs))

    def test_passing_preflight_keeps_translation_enabled(self) -> None:
        self.assertTrue(WorkflowController._preflight_translation(self._WorkingTranslator()))

    def test_no_translator_means_paused(self) -> None:
        self.assertFalse(WorkflowController._preflight_translation(None))


class TranslationStatusLabelTests(unittest.TestCase):
    def test_reason_string_is_stable(self) -> None:
        self.assertEqual(MESSAGE_TOO_MANY_GENES_REASON, "基因过多")
        self.assertEqual(TRANSLATION_STATUS_PAUSED, "paused")


class UndefinedNameGuardTests(unittest.TestCase):
    """静态兜底：捕获"用了不存在的变量名"这类本可避免的运行时错误。

    实际教训：metascape_client 里给文本框加 JS 聚焦时写了不存在的 `selector`，
    异常被兜底捕获后静默回退到慢速鼠标点击，直到真机日志才暴露。
    """

    @staticmethod
    def _module_files() -> list[Path]:
        return sorted((Path(__file__).resolve().parent.parent / "metabatch").glob("*.py"))

    def test_no_undefined_names_in_package(self) -> None:
        problems: list[str] = []
        for path in self._module_files():
            problems.extend(_check_module(path))
        self.assertEqual(problems, [])

    def test_guard_detects_the_real_regression(self) -> None:
        """自检：这个守卫必须能抓到当初那个 `selector is not defined`。"""
        with tempfile.TemporaryDirectory() as tmp:
            sample = Path(tmp) / "sample.py"
            sample.write_text(
                "def run(page, textbox_selector):\n"
                "    try:\n"
                "        page.evaluate('(s) => s', selector)\n"
                "    except Exception:\n"
                "        pass\n",
                encoding="utf-8",
            )
            problems = _check_module(sample)
        self.assertTrue(any("selector" in item for item in problems), problems)


class GeneLimitParsingTests(unittest.TestCase):
    """界面上可以填写自己的基因数上限。"""

    def test_empty_keeps_previous_value(self) -> None:
        from metabatch.gui import _parse_gene_limit

        self.assertEqual(_parse_gene_limit("", fallback=3000), 3000)
        self.assertEqual(_parse_gene_limit("   "), DEFAULT_MAX_GENES_PER_FILE)

    def test_valid_values_are_used(self) -> None:
        from metabatch.gui import _parse_gene_limit

        self.assertEqual(_parse_gene_limit("800"), 800)
        self.assertEqual(_parse_gene_limit(" 12000 "), 12000)

    def test_invalid_values_raise(self) -> None:
        from metabatch.gui import _parse_gene_limit

        with self.assertRaises(ValueError):
            _parse_gene_limit("abc")
        with self.assertRaises(ValueError):
            _parse_gene_limit("0")
        with self.assertRaises(ValueError):
            _parse_gene_limit("-5")


class GeneReportTests(unittest.TestCase):
    """结果清单：基因数 + 是否做了 Metascape + 是否翻译。"""

    def test_metascape_state_wording(self) -> None:
        self.assertEqual(_metascape_state(TaskStatus.SUCCESS, ""), "已做")
        self.assertEqual(_metascape_state(TaskStatus.PARTIAL, ""), "已做")
        self.assertEqual(_metascape_state(TaskStatus.SKIPPED, SKIPPED_REASON_EXISTING), "已有结果")
        self.assertEqual(_metascape_state(TaskStatus.SKIPPED, MESSAGE_TOO_MANY_GENES_REASON), "未做")
        self.assertEqual(_metascape_state(TaskStatus.FAILED, ""), "未做")
        # 中途停止不算"已做"（没有结果文件），标为部分完成
        self.assertEqual(_metascape_state(TaskStatus.STOPPED, ""), "部分完成")

    def test_translation_state_wording(self) -> None:
        self.assertEqual(_translation_state("translated", TaskStatus.SUCCESS), "已翻译")
        self.assertEqual(_translation_state("paused", TaskStatus.SUCCESS), "已暂停")
        self.assertEqual(_translation_state("failed", TaskStatus.PARTIAL), "翻译失败")
        self.assertEqual(_translation_state("partial", TaskStatus.SUCCESS), "未翻译")
        self.assertEqual(_translation_state("", TaskStatus.SKIPPED), "未知")

    def test_state_wording_follows_ui_language(self) -> None:
        set_ui_language(UI_LANGUAGE_EN)
        try:
            self.assertEqual(_metascape_state(TaskStatus.SUCCESS, ""), tr("已做"))
            self.assertEqual(_translation_state("translated", TaskStatus.SUCCESS), tr("已翻译"))
        finally:
            set_ui_language(UI_LANGUAGE_ZH)

    def test_report_csv_columns_and_encoding(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            input_dir = root / "genes"
            output_dir = root / "out"
            input_dir.mkdir()
            output_dir.mkdir()
            gene = input_dir / "a.txt"
            gene.write_text("TP53\nBRCA1", encoding="utf-8")
            records = [
                GeneFileRecord(
                    source_file=gene,
                    gene_count=2,
                    status=TaskStatus.SUCCESS,
                    translation_status="translated",
                    message="处理成功",
                ),
                GeneFileRecord(
                    source_file=input_dir / "b.txt",
                    gene_count=9000,
                    status=TaskStatus.SKIPPED,
                    skipped_reason=MESSAGE_TOO_MANY_GENES_REASON,
                    message="基因数 9000 超过上限 5000，已跳过该列表。",
                ),
            ]

            path = WorkflowController._write_gene_report(output_dir, records, input_dir)

            self.assertEqual(path.name, "metabatch_gene_report.csv")
            raw = path.read_bytes()
            self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "报告应为 UTF-8 BOM，便于 Excel 直接打开")
            with path.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(len(rows), 2)
            self.assertEqual(
                list(rows[0]),
                [
                    "source_file",
                    "relative_path",
                    "file_name",
                    "gene_count",
                    "metascape",
                    "translated",
                    "status",
                    "skipped_reason",
                    "translation_status",
                    "message",
                ],
            )
            self.assertEqual(rows[0]["relative_path"], "a.txt")
            self.assertEqual(rows[0]["gene_count"], "2")
            self.assertEqual(rows[0]["metascape"], "已做")
            self.assertEqual(rows[0]["translated"], "已翻译")
            self.assertEqual(rows[1]["gene_count"], "9000")
            self.assertEqual(rows[1]["metascape"], "未做")

    def test_summary_csv_contains_gene_count(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            output_dir = Path(tmp)
            results = [
                TaskResult(
                    source_file=Path("a.txt"),
                    status=TaskStatus.SUCCESS,
                    download_path=None,
                    translated_path=None,
                    message="处理成功",
                    elapsed_seconds=1.0,
                    translation_status="translated",
                    gene_count=42,
                )
            ]

            WorkflowController._write_summary(output_dir, results)

            with (output_dir / "metabatch_summary.csv").open(encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(rows[0]["gene_count"], "42")
            self.assertEqual(rows[0]["translation_status"], "translated")


class _FakeMetascapeClient:
    """替身：不开浏览器，直接写一个合法的结果 xlsx，并记录收到的代理参数。"""

    proxy_urls: list[str | None] = []

    def __init__(self, headless: bool = True, proxy_url: str | None = None) -> None:
        type(self).proxy_urls.append(proxy_url)

    def __enter__(self) -> "_FakeMetascapeClient":
        return self

    def __exit__(self, *exc) -> None:
        return None

    def process_gene_list(self, download_path: Path, **_kwargs):
        from metabatch.models import DownloadedArtifacts

        write_result_workbook(download_path, ["Category", "Description"])
        return DownloadedArtifacts(download_path=download_path, excel_path=download_path)


class WorkflowRunSmokeTests(unittest.TestCase):
    """端到端走一遍 run() → _process_single_file()，抓"真机一跑就全部失败"的回归。

    实际教训：_process_single_file() 用了只在 run() 里定义的 metascape_connection_mode，
    76 个文件 0.2 秒内全部 NameError 失败，而原有单元测试都没走到这条路径。
    """

    def _run(self, root: Path, **kwargs) -> tuple[list[TaskResult], Path]:
        from unittest import mock
        from threading import Event

        input_dir = root / "genes"
        output_dir = root / "out"
        write_gene_file(input_dir / "a.txt", 5)
        write_gene_file(input_dir / "b.txt", 5)
        config = AppConfig(
            input_dir=input_dir,
            output_dir=output_dir,
            enable_translation=False,
            concurrency=1,
            min_delay_seconds=0,
            max_delay_seconds=0,
            **kwargs,
        )
        _FakeMetascapeClient.proxy_urls = []
        with mock.patch("metabatch.workflow.MetascapeClient", _FakeMetascapeClient), mock.patch(
            "metabatch.workflow.make_log_path", return_value=root / "run.log"
        ):
            results = WorkflowController().run(config, Event(), lambda *_: None, lambda *_: None)
        return results, output_dir

    def test_run_processes_files_without_errors(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            results, output_dir = self._run(Path(tmp))
            self.assertEqual([item.status for item in results], [TaskStatus.SUCCESS, TaskStatus.SUCCESS], results)
            self.assertEqual(_FakeMetascapeClient.proxy_urls, [None, None])
            with (output_dir / "metabatch_gene_report.csv").open(encoding="utf-8-sig", newline="") as handle:
                self.assertEqual({row["status"] for row in csv.DictReader(handle)}, {"success"})

    def test_unreachable_proxy_falls_back_to_direct(self) -> None:
        from unittest import mock

        from metabatch.models import ConnectionMode

        with tempfile.TemporaryDirectory() as tmp, mock.patch(
            "metabatch.workflow.proxy_is_reachable", return_value=False
        ):
            results, _ = self._run(Path(tmp), metascape_connection_mode=ConnectionMode.PROXY, proxy_port=51888)
            self.assertTrue(all(item.status is TaskStatus.SUCCESS for item in results), results)
            self.assertEqual(_FakeMetascapeClient.proxy_urls, [None, None])

    def test_worker_crash_is_reported_as_failed_in_gene_report(self) -> None:
        from unittest import mock

        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(
            _FakeMetascapeClient, "process_gene_list", side_effect=RuntimeError("boom")
        ):
            results, output_dir = self._run(Path(tmp))
            self.assertTrue(all(item.status is TaskStatus.FAILED for item in results), results)
            with (output_dir / "metabatch_gene_report.csv").open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            # 之前占位是 SUCCESS，崩溃的文件会在清单里被算成"已做"
            self.assertEqual({row["status"] for row in rows}, {"failed"})
            self.assertEqual({row["metascape"] for row in rows}, {"未做"})


if __name__ == "__main__":
    unittest.main()
