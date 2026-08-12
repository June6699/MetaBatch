"""Batch-translate Metascape Enrichment descriptions in an existing result tree."""

from __future__ import annotations

import argparse
import csv
import json
import os
import time
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

from metabatch.excel_processor import ExcelProcessingError, finalize_workbook_layout


SYSTEM_PROMPT = (
    "Translate biomedical enrichment term descriptions from English to Simplified Chinese. "
    "Preserve gene symbols, pathway identifiers, abbreviations, and punctuation when needed. "
    "Return only a JSON object with one key, translations, whose value is a JSON array of "
    "translated strings in exactly the same order as the input items."
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results-dir", type=Path, required=True)
    parser.add_argument("--api-base-url", default="https://aibz.cc/v1")
    parser.add_argument("--model", default="deepseek-v4-flash")
    parser.add_argument("--batch-size", type=int, default=20)
    parser.add_argument("--summary-name", default="metabatch_translation_summary.tsv")
    parser.add_argument("--progress-log-name", default="metabatch_translation_run.log")
    return parser.parse_args()


def endpoint(base_url: str) -> str:
    base = base_url.rstrip("/")
    if base.endswith("/chat/completions"):
        return base
    return f"{base}/chat/completions"


def translation_api_key() -> str:
    """Read the credential without requiring a shell to expose it to a child process."""
    api_key = os.environ.get("METABATCH_TRANSLATION_API_KEY", "").strip()
    if api_key or os.name != "nt":
        return api_key

    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, "Environment") as key:
            value, _ = winreg.QueryValueEx(key, "METABATCH_TRANSLATION_API_KEY")
        return str(value).strip()
    except (FileNotFoundError, OSError):
        return ""


def workbook_state(path: Path) -> tuple[str, list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Enrichment" not in workbook.sheetnames:
            return "no_enrichment_sheet", []
        sheet = workbook["Enrichment"]
        headers = [str(cell.value).strip() for cell in next(sheet.iter_rows(max_row=1))]
        if "Description" not in headers:
            return "no_description_column", []
        if "中文描述" in headers:
            return "already_translated", []
        description_index = headers.index("Description")
        descriptions = [
            "" if row[description_index].value is None else str(row[description_index].value).strip()
            for row in sheet.iter_rows(min_row=2)
        ]
        return "ready", descriptions
    finally:
        workbook.close()


def parse_translations(content: str, expected_count: int) -> list[str]:
    value = content.strip()
    if value.startswith("```"):
        value = value.split("\n", 1)[1] if "\n" in value else ""
        value = value.rsplit("```", 1)[0].strip()
    payload: Any = json.loads(value)
    translations = payload.get("translations") if isinstance(payload, dict) else None
    if not isinstance(translations, list) or len(translations) != expected_count:
        raise ValueError("Translation API returned an invalid translation array.")
    return [str(item).strip() for item in translations]


def translate_batch(
    url: str,
    api_key: str,
    model: str,
    items: list[str],
) -> list[str]:
    payload = {
        "model": model,
        "temperature": 0,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": json.dumps({"items": items}, ensure_ascii=False)},
        ],
    }
    last_error: Exception | None = None
    for attempt in range(1, 4):
        try:
            response = requests.post(
                url,
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json=payload,
                timeout=60,
            )
            response.raise_for_status()
            body = response.json()
            content = body["choices"][0]["message"]["content"]
            return parse_translations(str(content), len(items))
        except (KeyError, ValueError, requests.RequestException) as exc:
            last_error = exc
            if attempt < 3:
                time.sleep(2**attempt)
    raise RuntimeError(f"Translation batch failed: {last_error}")


def write_translations(path: Path, translations: dict[str, str]) -> int:
    workbook = load_workbook(path)
    try:
        if "Enrichment" not in workbook.sheetnames:
            raise ExcelProcessingError("Results workbook has no Enrichment sheet.")
        sheet = workbook["Enrichment"]
        headers = [str(sheet.cell(row=1, column=index).value).strip() for index in range(1, sheet.max_column + 1)]
        if "Description" not in headers:
            raise ExcelProcessingError("Enrichment sheet has no Description column.")
        if "中文描述" in headers:
            raise ExcelProcessingError("Enrichment sheet is already translated.")
        description_column = headers.index("Description") + 1
        target_column = sheet.max_column + 1
        sheet.cell(row=1, column=target_column, value="中文描述")
        for row_index in range(2, sheet.max_row + 1):
            description = sheet.cell(row=row_index, column=description_column).value
            source = "" if description is None else str(description).strip()
            sheet.cell(row=row_index, column=target_column, value=translations.get(source, ""))
        workbook.save(path)
        return max(sheet.max_row - 1, 0)
    finally:
        workbook.close()


def append_progress(path: Path, line: str) -> None:
    with path.open("a", encoding="utf-8", newline="\n") as handle:
        handle.write(line + "\n")


def main() -> None:
    args = parse_args()
    results_dir = args.results_dir.resolve()
    api_key = translation_api_key()
    if not api_key:
        raise SystemExit("METABATCH_TRANSLATION_API_KEY is required.")
    if not results_dir.is_dir():
        raise SystemExit(f"Results directory does not exist: {results_dir}")
    if args.batch_size < 1:
        raise SystemExit("batch-size must be at least 1.")

    progress_log = results_dir / args.progress_log_name
    progress_log.write_text("", encoding="utf-8", newline="\n")
    workbooks = sorted(results_dir.rglob("*_metascape.xlsx"))
    records: list[dict[str, object]] = []
    ready: list[tuple[Path, list[str]]] = []
    unique_descriptions: set[str] = set()
    for path in workbooks:
        status, descriptions = workbook_state(path)
        record = {
            "file": path.relative_to(results_dir).as_posix(),
            "status": status,
            "description_rows": len(descriptions),
            "message": "",
        }
        records.append(record)
        if status == "ready":
            ready.append((path, descriptions))
            unique_descriptions.update(description for description in descriptions if description)

    translations: dict[str, str] = {}
    unique_items = sorted(unique_descriptions)
    api_url = endpoint(args.api_base_url)
    for start in range(0, len(unique_items), args.batch_size):
        chunk = unique_items[start : start + args.batch_size]
        translated = translate_batch(api_url, api_key, args.model, chunk)
        translations.update(dict(zip(chunk, translated, strict=True)))
        append_progress(progress_log, f"API {min(start + len(chunk), len(unique_items))}/{len(unique_items)} unique descriptions translated")
        time.sleep(0.1)

    record_by_file = {str(record["file"]): record for record in records}
    for index, (path, _) in enumerate(ready, start=1):
        record = record_by_file[path.relative_to(results_dir).as_posix()]
        try:
            record["description_rows"] = write_translations(path, translations)
            finalize_workbook_layout(path, sheet_name="Enrichment", auto_fit_headers=["Description", "中文描述"])
            record["status"] = "translated"
        except (ExcelProcessingError, OSError, ValueError) as exc:
            record["status"] = "failed"
            record["message"] = str(exc)
        append_progress(progress_log, f"WORKBOOK {index}/{len(ready)} {record['status']}: {record['file']}")

    summary_path = results_dir / args.summary_name
    with summary_path.open("w", encoding="utf-8", newline="\n") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["file", "status", "description_rows", "message"],
            delimiter="\t",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(records)

    counts: dict[str, int] = {}
    for record in records:
        status = str(record["status"])
        counts[status] = counts.get(status, 0) + 1
    append_progress(progress_log, f"SUMMARY {counts}")
    print(f"Summary: {counts}", flush=True)
    if counts.get("failed", 0):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
