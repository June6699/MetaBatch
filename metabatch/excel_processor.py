from __future__ import annotations

from pathlib import Path
from threading import Event
from typing import Callable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .i18n import tr
from .translator import Translator

RowProgressCallback = Callable[[int, int], None]

DEFAULT_TARGET_HEADER = "中文描述"


class ExcelProcessingError(Exception):
    """Raised when the downloaded workbook cannot be processed."""


def inspect_enrichment_sheet(workbook_path: Path, target_header: str = DEFAULT_TARGET_HEADER) -> int:
    """检查 xlsx 是否已有可用于翻译的 Enrichment 表。

    返回可翻译的数据行数；表中已有目标语言列时返回 ``0``（表示无需再翻译）。
    缺少 Enrichment 表或 Description 列时抛出 :class:`ExcelProcessingError`。
    """
    try:
        workbook = load_workbook(workbook_path, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExcelProcessingError(tr("结果文件已损坏，无法读取：{error}").format(error=exc)) from exc
    try:
        if "Enrichment" not in workbook.sheetnames:
            raise ExcelProcessingError(
                tr("结果文件中未找到 Enrichment 工作表，可能因为基因过少未产生富集结果。")
            )
        sheet = workbook["Enrichment"]
        _find_header_column(sheet, "Description")
        if _find_header_column(sheet, target_header, raise_if_missing=False) is not None:
            return 0
        return max(sheet.max_row - 1, 0)
    finally:
        workbook.close()


def translate_workbook(
    workbook_path: Path,
    translator: Translator,
    stop_event: Event | None = None,
    progress_callback: RowProgressCallback | None = None,
    target_header: str = DEFAULT_TARGET_HEADER,
) -> Path:
    workbook = load_workbook(workbook_path)
    try:
        # 校验 Enrichment 工作表存在
        if "Enrichment" not in workbook.sheetnames:
            raise ExcelProcessingError(
                tr("结果文件中未找到 Enrichment 工作表，可能因为基因过少未产生富集结果。")
            )
        sheet = workbook["Enrichment"]

        # 定位 Description 列，并确认目标列尚不存在（避免重复写入）
        description_column = _find_header_column(sheet, "Description")
        existing_column = _find_header_column(sheet, target_header, raise_if_missing=False)
        if existing_column is not None:
            raise ExcelProcessingError(
                tr("该文件已经存在“{column}”列，已跳过以避免重复写入。").format(column=target_header)
            )

        # 新增目标列并写入表头
        target_column = sheet.max_column + 1
        sheet.cell(row=1, column=target_column, value=target_header)

        # 收集全部 Description 作为待翻译源（空值统一为空串）
        total_rows = max(sheet.max_row - 1, 0)
        row_indexes = list(range(2, sheet.max_row + 1))
        descriptions = [sheet.cell(row=row_index, column=description_column).value for row_index in row_indexes]
        sources = ["" if value is None else str(value) for value in descriptions]

        # 一次性批量调用：是否整块发送由翻译器自行决定
        translated_values: list[str] = []
        translate_batch = getattr(translator, "translate_batch", None)
        if callable(translate_batch):
            def _forward_progress(done: int, total: int) -> None:
                if progress_callback:
                    progress_callback(done, total)

            batch_result = list(
                translate_batch(
                    sources,
                    progress_callback=_forward_progress,
                    stop_event=stop_event,
                )
            )
            if len(batch_result) != len(sources):
                raise ExcelProcessingError(tr("翻译接口返回的批量结果数量不匹配。"))
            translated_values.extend(batch_result)
        else:
            # 兼容只实现单条 translate 的轻量翻译器
            for index, source in enumerate(sources, start=1):
                if stop_event and stop_event.is_set():
                    raise ExcelProcessingError(tr("用户已停止任务。"))
                translated_values.append(translator.translate(source))
                if progress_callback:
                    progress_callback(index, total_rows)
        if progress_callback:
            progress_callback(total_rows, total_rows)

        # 按原顺序把翻译结果写回目标列
        for row_index, translated in zip(row_indexes, translated_values, strict=True):
            sheet.cell(row=row_index, column=target_column, value=translated)

        # 自动列宽并保存
        _finalize_workbook_layout(workbook, sheet_name="Enrichment", auto_fit_headers=["Description", target_header])
        workbook.save(workbook_path)
        return workbook_path
    finally:
        workbook.close()


def finalize_workbook_layout(
    workbook_path: Path,
    sheet_name: str = "Enrichment",
    auto_fit_headers: list[str] | None = None,
) -> Path:
    workbook = load_workbook(workbook_path)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ExcelProcessingError(
                tr("结果文件中未找到 Enrichment 工作表，可能因为基因过少未产生富集结果。")
            )
        headers = auto_fit_headers or [DEFAULT_TARGET_HEADER, "Description"]
        _finalize_workbook_layout(workbook, sheet_name=sheet_name, auto_fit_headers=headers)
        workbook.save(workbook_path)
        return workbook_path
    finally:
        workbook.close()


def _find_header_column(sheet, header_name: str, raise_if_missing: bool = True) -> int | None:
    for column_index in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=column_index).value
        if str(value).strip() == header_name:
            return column_index

    if raise_if_missing:
        raise ExcelProcessingError(tr("未找到列“{header}”。").format(header=header_name))
    return None


def _finalize_workbook_layout(workbook, sheet_name: str, auto_fit_headers: list[str]) -> None:
    # 定位目标表及其索引
    sheet = workbook[sheet_name]
    sheet_index = workbook.sheetnames.index(sheet_name)

    # 只让目标表的视图标记为选中
    for worksheet in workbook.worksheets:
        worksheet.sheet_view.tabSelected = worksheet.title == sheet_name

    # 同步工作簿级活动标签与首屏标签
    workbook.active = sheet_index
    if workbook.views:
        workbook_view = workbook.views[0]
        workbook_view.activeTab = sheet_index
        workbook_view.firstSheet = sheet_index

    # 调整目标列列宽
    _auto_fit_columns(sheet, auto_fit_headers)


def _auto_fit_columns(sheet, headers: list[str]) -> None:
    # 建立表头名到列号的映射
    header_to_index: dict[str, int] = {}
    for column_index in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=column_index).value
        header_to_index[str(value).strip()] = column_index

    # 对每个目标列，按表头与全部数据的显示宽度求最大值
    for header in headers:
        column_index = header_to_index.get(header)
        if column_index is None:
            continue

        max_length = _display_text_width(header)
        for row_index in range(2, sheet.max_row + 1):
            value = sheet.cell(row=row_index, column=column_index).value
            text = "" if value is None else str(value)
            max_length = max(max_length, _display_text_width(text))

        # 宽度夹在 12~80，标记 bestFit
        column_letter = get_column_letter(column_index)
        dimension = sheet.column_dimensions[column_letter]
        dimension.bestFit = True
        dimension.width = min(max(max_length + 2, 12), 80)


def _display_text_width(text: str) -> int:
    width = 0
    for char in text:
        width += 2 if ord(char) > 127 else 1
    return width
