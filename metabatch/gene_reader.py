from __future__ import annotations

import csv
import fnmatch
import re
from io import StringIO
from pathlib import Path
from typing import Iterable, Sequence

import xlrd
from openpyxl import load_workbook

from .i18n import tr
from .paths import is_generated_result_path

SUPPORTED_INPUT_SUFFIXES = {".txt", ".csv", ".tsv", ".xlsx", ".xls"}

# 超限文案：gene_reader 提供、workflow 复用，避免两处措辞不一致
MESSAGE_TOO_MANY_GENES = "基因数 {count} 超过上限 {limit}，已跳过该列表。"
# 这两个是写进 summary / TaskResult 的"原因键"，始终保留中文原始值（不参与界面翻译），
# 否则英文界面下按字符串比较统计会失配。
MESSAGE_TOO_MANY_GENES_REASON = "基因过多"
SKIPPED_REASON_EXISTING = "已有结果"
SKIPPED_REASON_EXISTING_MESSAGE = "已存在结果文件：{path}"


class GeneReadError(Exception):
    """Raised when a gene file cannot be interpreted."""


def normalize_patterns(patterns: str | Sequence[str] | None) -> list[str]:
    """Parse user-supplied glob patterns (``;``/``,`` separated) into lowercase tokens."""
    if patterns is None:
        return []

    # 字符串按 ; , 拆成多个模式；序列直接逐项处理
    raw: Iterable[str] = (patterns,) if isinstance(patterns, str) else patterns
    result: list[str] = []
    for item in raw:
        for token in re.split(r"[;,]", str(item)):
            token = token.strip()
            if token:
                result.append(token.lower())
    return result


def discover_gene_files(
    input_dir: Path,
    patterns: str | Sequence[str] | None = None,
) -> list[Path]:
    # 解析用户模式；为空表示不过滤
    compiled_patterns = normalize_patterns(patterns)

    # 递归扫描：只收文件、后缀受支持、排除已生成的结果、模式匹配通过
    matched: list[Path] = []
    for path in input_dir.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_INPUT_SUFFIXES:
            continue
        if _looks_like_generated_result(path):
            continue
        if compiled_patterns and not _matches_patterns(input_dir, path, compiled_patterns):
            continue
        matched.append(path)

    # 按相对路径（不区分大小写）排序，保证每次运行顺序稳定
    return sorted(
        matched,
        key=lambda item: str(item.relative_to(input_dir)).lower(),
    )


def _matches_patterns(input_dir: Path, path: Path, compiled_patterns: list[str]) -> bool:
    # 模式含分隔符时匹配相对路径，否则只匹配文件名；统一不区分大小写
    relative_name = str(path.relative_to(input_dir)).replace("\\", "/").lower()
    file_name = path.name.lower()

    for pattern in compiled_patterns:
        normalized = pattern.replace("\\", "/")
        target = relative_name if "/" in normalized else file_name
        if fnmatch.fnmatch(target, normalized):
            return True
    return False


def read_gene_list(path: Path, column_selector: str | None = None) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return _normalize_values(_decode_text(path).splitlines())
    if suffix in {".csv", ".tsv"}:
        delimiter = "," if suffix == ".csv" else "\t"
        return _read_delimited(path, delimiter, column_selector)
    if suffix == ".xlsx":
        return _read_xlsx(path, column_selector)
    if suffix == ".xls":
        return _read_xls(path, column_selector)
    raise GeneReadError(tr("不支持的输入格式：{suffix}").format(suffix=path.suffix))


def count_gene_list(path: Path, column_selector: str | None = None) -> int:
    """只统计基因条数，供运行前预扫描与阈值判定使用。"""
    return len(read_gene_list(path, column_selector))


def exceeds_gene_limit(count: int, limit: int | None) -> bool:
    """上限为 None（未启用或无有效数值）时视为不限制。"""
    return limit is not None and count > limit


def too_many_genes_message(count: int, limit: int) -> str:
    return tr(MESSAGE_TOO_MANY_GENES).format(count=count, limit=limit)


def _normalize_values(values: Iterable[object]) -> list[str]:
    genes: list[str] = []
    for value in values:
        text = str(value).strip()
        if text:
            genes.append(text)
    return genes


def _decode_text(path: Path) -> str:
    # Windows 下常见 UTF-8 BOM 与 GBK/GB2312 编码文件，utf-8-sig 兼容前两者之外的情况交给 gb18030 兜底。
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise GeneReadError(tr("无法解码文件内容，请将文件另存为 UTF-8 编码：{name}").format(name=path.name))


def _looks_like_generated_result(path: Path) -> bool:
    # 名称命中：解压目录、metascape_result 或 _metascape 结果后缀
    lowered = path.stem.lower()
    if lowered.endswith("_extracted") or lowered.startswith("metascape_result.") or is_generated_result_path(path):
        return True

    # 内容命中：含 Enrichment 工作表的 Excel 是 Metascape 输出，不是基因列表
    if path.suffix.lower() in {".xlsx", ".xls"}:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                return "Enrichment" in workbook.sheetnames
            finally:
                workbook.close()
        except Exception:
            return False

    return False


def _resolve_column_index(
    selector: str | None,
    header_row: Sequence[object] | None = None,
) -> tuple[int, bool]:
    # 未指定列：默认取第 1 列，不跳过表头
    if not selector:
        return 0, False
    token = selector.strip()

    # 先尝试按表头名匹配（需要可读表头）
    if header_row is not None:
        normalized_headers = [str(cell).strip() for cell in header_row]
        try:
            return normalized_headers.index(token), True
        except ValueError:
            pass

    # 纯数字：1-based 列号
    if token.isdigit():
        index = int(token) - 1
        if index < 0:
            raise GeneReadError(tr("列号必须从 1 开始。"))
        return index, False

    # 纯字母：Excel 列字母（A、B、AA…），换算成 0-based
    if re.fullmatch(r"[A-Za-z]+", token):
        excel_index = 0
        for char in token.upper():
            excel_index = excel_index * 26 + (ord(char) - ord("A") + 1)
        return excel_index - 1, False

    # 走到这里说明给的是表头名但未匹配上
    if header_row is None:
        raise GeneReadError(tr("按表头名选择列时需要可读取表头的结构化文件。"))
    normalized_headers = [str(cell).strip() for cell in header_row]
    try:
        return normalized_headers.index(token), True
    except ValueError as exc:
        raise GeneReadError(tr("未找到名为“{name}”的列。").format(name=token)) from exc


def _read_delimited(path: Path, delimiter: str, column_selector: str | None) -> list[str]:
    # 读取并按分隔符解析全部行
    rows = list(csv.reader(StringIO(_decode_text(path)), delimiter=delimiter))
    if not rows:
        return []

    # 解析目标列，决定是否跳过表头
    header_row = rows[0]
    column_index, skip_header = _resolve_column_index(column_selector, header_row)
    start_index = 1 if skip_header else 0

    # 逐行取目标列（缺列的行跳过）
    values: list[str] = []
    for row in rows[start_index:]:
        if column_index >= len(row):
            continue
        values.append(row[column_index])
    return _normalize_values(values)


def _read_xlsx(path: Path, column_selector: str | None) -> list[str]:
    # 打开工作簿，读取第一个工作表全部行
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []

    # 解析目标列，决定是否跳过表头
    header_row = list(rows[0])
    column_index, skip_header = _resolve_column_index(column_selector, header_row)
    start_index = 1 if skip_header else 0

    # 逐行取目标列
    values: list[object] = []
    for row in rows[start_index:]:
        if column_index >= len(row):
            continue
        values.append(row[column_index])
    return _normalize_values(values)


def _read_xls(path: Path, column_selector: str | None) -> list[str]:
    # 打开老式 xls，读取第一个工作表
    workbook = xlrd.open_workbook(path.as_posix())
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        return []

    # 解析目标列，决定是否跳过表头
    header_row = sheet.row_values(0)
    column_index, skip_header = _resolve_column_index(column_selector, header_row)
    start_index = 1 if skip_header else 0

    # 逐行取目标列
    values: list[object] = []
    for row_index in range(start_index, sheet.nrows):
        row_values = sheet.row_values(row_index)
        if column_index >= len(row_values):
            continue
        values.append(row_values[column_index])
    return _normalize_values(values)
